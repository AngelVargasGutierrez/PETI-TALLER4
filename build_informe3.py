# -*- coding: utf-8 -*-
"""Etapa 3: contenido de las secciones 1 a 7."""
import copy
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

PATH = "SI886-LAB-04_vargas.docx"
REPO = "https://github.com/AngelVargasGutierrez/PETI-TALLER4"
GRAF = "graficos/S04_diagnostico.png"
ANEXO_A = "evidencias/S04/anexo_A_crehana_pagina.png"

d = docx.Document(PATH)


def find_para(text, style=None, contains=False):
    for para in d.paragraphs:
        t = para.text.strip()
        ok = (text in t) if contains else (t == text)
        if ok and (style is None or para.style.name == style):
            return para
    raise ValueError(f"No se encontro parrafo: {text!r}")


def insert_paragraph_after(paragraph, text="", style=None):
    new_p = copy.deepcopy(paragraph._p)
    for child in list(new_p):
        new_p.remove(child)
    paragraph._p.addnext(new_p)
    from docx.text.paragraph import Paragraph
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def add_paragraphs(after_para, items, style=None):
    cur = after_para
    for it in items:
        cur = insert_paragraph_after(cur, it, style=style)
    return cur


def add_bullets(after_para, items):
    return add_paragraphs(after_para, [f"•  {it}" for it in items], style="Compact")


def add_heading2_after(after_para, text):
    return insert_paragraph_after(after_para, text, style="Heading 2")


def add_image_paragraph(after_para, img_path, width_in=5.5, caption=None):
    cur = insert_paragraph_after(after_para, "")
    cur.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cur.add_run()
    run.add_picture(img_path, width=Inches(width_in))
    if caption:
        cur2 = insert_paragraph_after(cur, caption)
        cur2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in cur2.runs:
            r.italic = True
            r.font.size = Pt(9)
        return cur2
    return cur


def add_link_paragraph(after_para, label, url):
    cur = insert_paragraph_after(after_para, f"{label}: {url}")
    return cur


# =====================================================================
# SECCION 1
# =====================================================================
h = find_para("1.1 Título del evento práctico", style="Heading 2")
add_paragraphs(h, [
    "Taller de laboratorio 04 · Formulación y validación de la misión y la visión "
    "(SI-886 Planeamiento Estratégico de TI, sesión 2 en laboratorio, 100 minutos)."
])

h = find_para("1.2 Objetivos", style="Heading 2")
add_bullets(h, [
    "Elegir una empresa real del sector tecnológico y registrar su misión y su visión "
    "literales, con su fuente.",
    "Evaluar la misión con los cinco componentes, los siete defectos y las tres pruebas "
    "de calidad.",
    "Evaluar la visión con los cinco atributos y extraer sus métricas implícitas.",
    "Derivar la visión de la función de TI.",
    "Redactar las secciones 2.1 y 2.2 del PETI.",
])

h = find_para("1.3 Tiempo de duración", style="Heading 2")
add_paragraphs(h, [
    "100 minutos de sesión de laboratorio, distribuidos en seis pasos: Paso A — Elegir "
    "la empresa y registrar sus declaraciones (15 min); Paso B — Evaluar la misión "
    "(20 min); Paso C — Evaluar la visión (15 min); Paso D — Derivar la visión de TI y "
    "redactar las Secciones 2.1 y 2.2 (10 min); Paso E — Validar y corregir (25 min); "
    "Paso F — Registrar la evidencia y cerrar (15 min)."
])

h = find_para("1.4 Resultados de aprendizaje", style="Heading 2")
add_paragraphs(h, [
    "La guía de laboratorio (3-TALLER.md) no incluye un listado propio de resultados de "
    "aprendizaje: estos se definen en el sílabo del curso, documento que no formaba "
    "parte de la carpeta de la semana compartida para este taller. A partir de los "
    "objetivos de la sesión (Sección 1.2) y del criterio de éxito declarado en el reto, "
    "se derivan los siguientes resultados de aprendizaje esperados:",
])
add_bullets(h, [
    "Aplica un instrumento técnico y no una impresión subjetiva para diagnosticar la "
    "calidad de una declaración de misión institucional.",
    "Distingue una declaración de misión o visión genuina de un eslogan publicitario, "
    "utilizando pruebas de sustitución, decisión y reconocimiento.",
    "Deriva capacidades de tecnologías de información verificables — con estado actual "
    "y estado objetivo — a partir de elementos concretos de la identidad estratégica "
    "de una empresa real.",
])

# --- 1.5 Recursos: llenar tabla existente y agregar fila de draw.io ---
tabla_recursos = d.tables[0]
filas = tabla_recursos.rows
datos_recursos = [
    ["Navegador con acceso a internet", "Microsoft Edge (build del sistema)",
     "Consultar el sitio oficial de Crehana (misión, visión y datos de la empresa) y "
     "capturar la evidencia del Anexo A."],
    ["Python 3.11+ con matplotlib", "Python 3.13.7 / matplotlib 3.11.1",
     "Ejecutar MV01_diagnostico_declaraciones.py: aplicar los instrumentos de la teoría "
     "y generar el gráfico de diagnóstico."],
    ["LibreOffice Writer / Microsoft Word", "Microsoft Word (Microsoft 365)",
     "Redacción de las Secciones 2.1 y 2.2 del PETI, y de este informe."],
]
for i, fila_datos in enumerate(datos_recursos):
    row = filas[i + 1]
    for c, val in enumerate(fila_datos):
        row.cells[c].text = val

fila_extra = tabla_recursos.add_row()
fila_extra.cells[0].text = "draw.io"
fila_extra.cells[1].text = "No utilizado directamente"
fila_extra.cells[2].text = (
    "La Sección 2.2.4 exige explícitamente el resultado en formato de tabla de "
    "derivación (elemento de la visión → capacidad de negocio → capacidad de TI → "
    "estado actual → estado objetivo); dicha tabla se construyó directamente en "
    "Markdown y se documenta en el cuerpo del informe, por lo que no fue necesario "
    "un diagrama adicional en draw.io."
)

h = find_para("1.6 Seguridad", style="Heading 2")
add_bullets(h, [
    "La misión y la declaración de impacto de Crehana se copiaron literalmente, con la "
    "dirección de la página (https://www.crehana.com/sobre-nosotros/) y la fecha de "
    "consulta (10/09/2026).",
    "Son textos publicados por la propia Crehana; se citan entre comillas y se le "
    "atribuyen a ella en todo momento.",
    "Solo se consultaron páginas públicas de la empresa. Se intentó una verificación "
    "cruzada en el perfil oficial de LinkedIn de Crehana, que no fue accesible por "
    "requerir inicio de sesión; no se insistió con credenciales de ningún tipo.",
    "Todas las declaraciones originales revisadas estaban en español; no fue necesario "
    "traducir ningún texto.",
])

d.save(PATH)
print("Seccion 1 OK")
