# -*- coding: utf-8 -*-
"""Arma el informe SI886-LAB-04_vargas.docx a partir de la plantilla, con el
contenido real del Taller 04 (empresa: Crehana)."""
import copy
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

PATH = "SI886-LAB-04_vargas.docx"
REPO = "https://github.com/AngelVargasGutierrez/PETI-TALLER4"
GRAF = "graficos/S04_diagnostico.png"
ANEXO_A = "evidencias/S04/anexo_A_crehana_pagina.png"

d = docx.Document(PATH)


def insert_paragraph_after(paragraph, text="", style=None):
    new_p = copy.deepcopy(paragraph._p)
    for child in list(new_p):
        new_p.remove(child)
    paragraph._p.addnext(new_p)
    from docx.text.paragraph import Paragraph
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def set_paragraph_text(paragraph, text, bold=None):
    for r in list(paragraph.runs):
        r._element.getparent().remove(r._element)
    run = paragraph.add_run(text)
    if bold is not None:
        run.bold = bold
    return paragraph


def add_bullets(after_para, items, style="List Bullet"):
    cur = after_para
    for it in items:
        cur = insert_paragraph_after(cur, it, style=style)
    return cur


def add_paragraphs(after_para, items, style=None):
    cur = after_para
    for it in items:
        cur = insert_paragraph_after(cur, it, style=style)
    return cur


def delete_paragraph(paragraph):
    p = paragraph._p
    p.getparent().remove(p)


def add_image_paragraph(after_para, img_path, width_in=5.5, caption=None):
    cur = insert_paragraph_after(after_para, "")
    cur.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cur.add_run()
    run.add_picture(img_path, width=Inches(width_in))
    if caption:
        cur2 = insert_paragraph_after(cur, caption)
        cur2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in cur2.runs:
            r.italic = True
            r.font.size = Pt(9)
        return cur2
    return cur


# =====================================================================
# CARATULA
# =====================================================================
p = d.paragraphs
set_paragraph_text(p[11], "“Formulación y validación de la misión y la visión”", bold=True)
set_paragraph_text(p[14], "“SI-886 Planeamiento Estratégico de TI”", bold=True)
set_paragraph_text(p[17], "VARGAS GUTIERREZ, ANGEL JOSE · 20200669922")

d.save(PATH)
print("Etapa 1 (caratula) OK")
