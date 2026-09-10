# -*- coding: utf-8 -*-
"""Etapa 4: Seccion 2 (Paso A-F)."""
import copy
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

PATH = "SI886-LAB-04_vargas.docx"
REPO = "https://github.com/AngelVargasGutierrez/PETI-TALLER4"
GRAF = "graficos/S04_diagnostico.png"
ANEXO_A = "evidencias/S04/anexo_A_crehana_pagina.png"

d = docx.Document(PATH)


def find_para(text, style=None, contains=False):
    for para in d.paragraphs:
        t = para.text.strip()
        ok = (text in t) if contains else (t == text)
        if ok and (style is None or para.style.name == style):
            return para
    raise ValueError(f"No se encontro parrafo: {text!r}")


def insert_paragraph_after(paragraph, text="", style=None):
    new_p = copy.deepcopy(paragraph._p)
    for child in list(new_p):
        new_p.remove(child)
    paragraph._p.addnext(new_p)
    from docx.text.paragraph import Paragraph
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def add_paragraphs(after_para, items, style=None):
    cur = after_para
    for it in items:
        cur = insert_paragraph_after(cur, it, style=style)
    return cur


def add_bullets(after_para, items):
    return add_paragraphs(after_para, [f"•  {it}" for it in items], style="Compact")


def add_heading2_after(after_para, text):
    return insert_paragraph_after(after_para, text, style="Heading 2")


def add_image_paragraph(after_para, img_path, width_in=5.5, caption=None):
    cur = insert_paragraph_after(after_para, "")
    cur.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cur.add_run()
    run.add_picture(img_path, width=Inches(width_in))
    if caption:
        cur2 = insert_paragraph_after(cur, caption)
        cur2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in cur2.runs:
            r.italic = True
            r.font.size = Pt(9)
        return cur2
    return cur


# =====================================================================
# SECCION 2 -- PROCEDIMIENTO
# =====================================================================
h_pasoA = find_para("Paso A", style="Heading 2")
cur = add_paragraphs(h_pasoA, [
    "Qué se buscaba: identificar una empresa real del sector tecnológico y registrar, "
    "de forma literal, su misión y su visión, con la dirección de la página y la fecha "
    "de consulta.",
    "Acción: se eligió Crehana (plataforma de gestión y desarrollo de talento para "
    "equipos de Recursos Humanos, fundada en Lima, Perú, en 2015). Se consultó su "
    "página oficial «Sobre nosotros» (https://www.crehana.com/sobre-nosotros/, "
    "10/09/2026) y se transcribió literalmente el bloque titulado «Misión». Se "
    "verificó que la página no contiene ningún bloque titulado «Visión».",
])
add_image_paragraph(cur, ANEXO_A, width_in=5.5,
                     caption="Anexo A. Página oficial de Crehana consultada — "
                             "https://www.crehana.com/sobre-nosotros/, 10/09/2026.")
cur = find_para("Anexo A. Página oficial de Crehana consultada — "
                 "https://www.crehana.com/sobre-nosotros/, 10/09/2026.")
cur = add_paragraphs(cur, [
    f"Evidencia: {REPO}/blob/main/02_identidad/MV01_declaraciones.md"
])

h_pasoB = find_para("Paso B", style="Heading 2")
cur = add_paragraphs(h_pasoB, [
    "Qué se buscaba: evaluar la misión vigente con los cinco componentes, los siete "
    "defectos técnicos y las tres pruebas de calidad de la teoría.",
    "Acción: se escribió MV01_diagnostico_declaraciones.py con los datos reales de "
    "Crehana (el script original del curso no estaba disponible en la carpeta "
    "compartida de la semana; se reconstruyó aplicando literalmente los criterios de "
    "3-TALLER.md) y se ejecutó con Python 3.13.7 y matplotlib 3.11.1.",
])
cur = add_paragraphs(cur, [
    "Resultado: 1 de 5 componentes presente, 7 de 7 defectos identificados, 0 de 3 "
    "pruebas de calidad superadas. Veredicto: SE REFORMULA (porta 3 componentes o "
    "menos)."
])
add_image_paragraph(cur, GRAF, width_in=5.5,
                     caption="Figura 1. Gráfico resumen del diagnóstico de la misión de Crehana.")
cur = find_para("Figura 1. Gráfico resumen del diagnóstico de la misión de Crehana.")
cur = add_paragraphs(cur, [
    f"Evidencia: {REPO}/blob/main/evidencias/S04/diagnostico.txt (salida completa del "
    f"script) y {REPO}/blob/main/02_identidad/2.1_mision.md (diagnóstico redactado)."
])

h_pasoC = find_para("Paso C", style="Heading 2")
cur = add_paragraphs(h_pasoC, [
    "Qué se buscaba: evaluar la visión vigente con los cinco atributos de la teoría y "
    "extraer sus métricas implícitas.",
    "Acción: se revisó exhaustivamente la página oficial de Crehana; no se encontró "
    "ningún bloque titulado «Visión». Conforme a la instrucción del taller, esta "
    "ausencia se registró como hallazgo, y los cinco atributos se evaluaron sobre la "
    "declaración de impacto disponible («Impulsamos equipos. Potenciamos culturas. "
    "Transformamos resultados.»), dejando explícito que se trata de un sustituto y no "
    "de una visión declarada.",
])
cur = add_paragraphs(cur, [
    "Resultado: 0 de 5 atributos cumplidos (uno no evaluable por ausencia de meta "
    "declarada). Se documentaron, en su lugar, las cifras reales que sí publica la "
    "empresa (más de 1200 empresas cliente en LATAM, 10 años de trayectoria) como "
    "línea base disponible para cualquier meta futura.",
    f"Evidencia: {REPO}/blob/main/02_identidad/2.2_vision.md",
])

h_pasoD = add_heading2_after(cur, "Paso D")
cur = add_paragraphs(h_pasoD, [
    "Qué se buscaba: derivar la visión de la función de TI a partir de los elementos "
    "concretos de la identidad estratégica de la empresa, y redactar las Secciones 2.1 "
    "y 2.2 del PETI.",
    "Acción: se tomaron los tres elementos concretos y verificables que sí declara "
    "Crehana en su página oficial (la unificación del recorrido del talento, la escala "
    "de más de 1200 empresas cliente en LATAM y la declaración de impacto sobre "
    "cultura y resultados) y, para cada uno, se derivó la capacidad de negocio y la "
    "capacidad de TI que la habilita, con su estado actual y su estado objetivo.",
    "Resultado: visión de la función de TI redactada (Sección 2.2.3) y tabla de "
    "derivación de tres filas (Sección 2.2.4), cada una trazable a un fragmento "
    "textual concreto de la página oficial de la empresa.",
    f"Evidencia: {REPO}/blob/main/02_identidad/2.2_vision.md",
])

h_pasoE = add_heading2_after(cur, "Paso E")
cur = add_paragraphs(h_pasoE, [
    "Qué se buscaba: validar el resultado con tres comprobaciones y corregir lo que "
    "fallara antes de cerrar la sesión.",
])
cur = add_bullets(cur, [
    "Prueba de sustitución con un competidor real (Buk): la variante «La misión de "
    "Buk es construir equipos listos para el futuro» sigue sonando igual de creíble "
    "que el original — la declaración no dice nada propio de Crehana. Confirma el "
    "veredicto del Paso B.",
    "Línea base de cada métrica de la visión, con su fuente: se comprobó que, al no "
    "existir una visión con metas propias, no hay ninguna cifra «implícita en la "
    "visión» que documentar; se registró honestamente esta ausencia en la Sección "
    "2.2.2, en vez de inventar una meta que la empresa no ha declarado.",
    "Cada fila de la tabla de derivación (Sección 2.2.4) se verificó contra el texto "
    "original de la página de Crehana: las tres provienen de fragmentos citados "
    "literalmente («unifica lo que antes estaba desconectado», «+1200 empresas», "
    "«Impulsamos equipos. Potenciamos culturas. Transformamos resultados.»), no de "
    "una idea del equipo sin respaldo textual.",
])
cur = add_paragraphs(cur, [
    "Las tres comprobaciones se cumplieron correctamente en la primera ejecución; no "
    "fue necesario corregir ningún resultado antes de cerrar la sesión.",
])

h_pasoF = add_heading2_after(cur, "Paso F")
cur = add_paragraphs(h_pasoF, [
    "Qué se buscaba: versionar lo producido, anotar la URL de cada resultado y "
    "responder la pregunta de transferencia.",
    f"Acción: se creó el repositorio dedicado {REPO}, con la rama taller-04 fusionada "
    f"a main mediante Pull Request ({REPO}/pull/1), y se aplicaron las etiquetas "
    f"v0.4 y taller-04 sobre el commit de cierre. El repositorio se referenció además "
    f"desde el repositorio acumulativo del curso "
    f"(https://github.com/AngelVargasGutierrez/peti-si886).",
])
cur = add_paragraphs(cur, [
    "Pregunta de transferencia — ¿qué riesgo correría una organización real si esto se "
    "hiciera mal? Si Crehana invirtiera en TI a partir de una misión y una visión "
    "genéricas y sin verificar, el portafolio de proyectos terminaría financiando "
    "capacidades que no se distinguen de las de cualquier competidor, sin ninguna "
    "métrica que permita comprobar si la inversión realmente acercó a la empresa al "
    "futuro que promete. La ausencia de una visión con metas verificables, además, "
    "deja a la función de TI sin un criterio objetivo para decidir qué proyecto "
    "priorizar primero."
])

d.save(PATH)
print("Seccion 2 (Pasos A-F) OK")
